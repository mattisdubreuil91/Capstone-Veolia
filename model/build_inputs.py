"""Step 1 - Inputs workbook: Veolia financial data, every number with its document and PDF page.

Run:  python build_inputs.py <folder with the 4 PDFs>
1. extract_statements.py reads the statement tables straight from the PDFs (no retyping).
2. Note data (debt, liquidity, deals, guidance) is listed below with its page.
3. A Checks block re-adds every statement with Excel formulas so a misread number shows up.
"""
import csv, subprocess, sys, pathlib, tempfile
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

PDF_DIR = sys.argv[1]
OUT = "Veolia_inputs.xlsx"
HERE = pathlib.Path(__file__).parent
F = "Arial"
BLUE, BLACK, GREY = "0000FF", "000000", "595959"
NAVY = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
NUM = '#,##0;(#,##0);"-"'
DOCS = {
    "FY25": "Veolia consolidated financial statements at 31/12/2025",
    "H126": "Veolia half-year report at 30/06/2026 (amendment to the 2025 URD)",
    "H124": "Veolia half-year report at 30/06/2024 (amendment to the 2023 URD)",
    "BRF": "Capstone project briefing",
}

# ------------------------------------------------------------------ 1. extract statements
tmp = pathlib.Path(tempfile.mkdtemp()) / "raw.csv"
subprocess.run([sys.executable, "-I", str(HERE / "extract_statements.py"), PDF_DIR, str(tmp)], check=True)
RAW = OrderedDict()   # (doc, section) -> OrderedDict(label -> {period: (value, page)})
for r in csv.DictReader(open(tmp)):
    key = (r["document"], r["section"])
    RAW.setdefault(key, OrderedDict()).setdefault(r["label"], {})[r["period"]] = (float(r["value"]), int(r["pdf_page"]))

def merge(parts):
    """parts = [(doc, section, [periods], doc_code)] -> ordered label list with {period: (value, src)}"""
    rows = OrderedDict()
    for doc, sec, periods, code in parts:
        for label, vals in RAW[(doc, sec)].items():
            label = label.replace("for the year", "for the period")
            for p in periods:
                if p in vals:
                    v, pg = vals[p]
                    rows.setdefault(label, {})[p] = (v, f"{code} p.{pg}")
    return rows

STATEMENTS = [
    ("Income statement", ["FY2024", "FY2025", "H1 2025", "H1 2026"],
     [("FY2025 statements", "Income statement", ["FY2024", "FY2025"], "FY25"),
      ("H1-2026 report", "Income statement", ["H1 2025", "H1 2026"], "H126")]),
    ("Balance sheet - assets", ["31/12/2024", "31/12/2025", "30/06/2026"],
     [("FY2025 statements", "Balance sheet - assets", ["31/12/2024", "31/12/2025"], "FY25"),
      ("H1-2026 report", "Balance sheet - assets", ["30/06/2026"], "H126")]),
    ("Balance sheet - equity and liabilities", ["31/12/2024", "31/12/2025", "30/06/2026"],
     [("FY2025 statements", "Balance sheet - equity and liabilities", ["31/12/2024", "31/12/2025"], "FY25"),
      ("H1-2026 report", "Balance sheet - equity and liabilities", ["30/06/2026"], "H126")]),
    ("Cash flow statement", ["FY2024", "FY2025", "H1 2025", "H1 2026"],
     [("FY2025 statements", "Cash flow statement", ["FY2024", "FY2025"], "FY25"),
      ("H1-2026 report", "Cash flow statement", ["H1 2025", "H1 2026"], "H126")]),
    ("Key figures (Veolia alternative performance measures, as published)",
     ["31/12/2022", "30/06/2023", "31/12/2023", "30/06/2024", "31/12/2024", "30/06/2025", "31/12/2025", "30/06/2026"],
     [("H1-2024 report", "Key figures", ["31/12/2022", "30/06/2023", "31/12/2023", "30/06/2024"], "H124"),
      ("H1-2026 report", "Key figures", ["31/12/2024", "30/06/2025", "31/12/2025", "30/06/2026"], "H126")]),
    ("Net financial debt bridge (Veolia definition, half-years)", ["H1 2023", "H1 2024", "H1 2025", "H1 2026"],
     [("H1-2024 report", "Net debt bridge", ["H1 2023", "H1 2024"], "H124"),
      ("H1-2026 report", "Net debt bridge", ["H1 2025", "H1 2026"], "H126")]),
]

# ------------------------------------------------------------------ 2. note data (from the notes and the management report)
NOTES = [
    ("Net financial debt - composition", ["30/06/2025", "31/12/2025", "30/06/2026"], [
        ("Non-current financial liabilities", {"30/06/2025": 20809, "30/06/2026": 22406}, "H126 p.28", ""),
        ("Current financial liabilities", {"30/06/2025": 9116, "30/06/2026": 11283}, "H126 p.28", ""),
        ("Bank overdrafts and other cash position items", {"30/06/2025": 156, "30/06/2026": 90}, "H126 p.28", ""),
        ("Bond issues", {"31/12/2025": 17992}, "FY25 p.66", ""),
        ("Other liabilities and bank overdrafts", {"31/12/2025": 9543}, "FY25 p.66", ""),
        ("IFRS 16 lease debt", {"31/12/2025": 1983}, "FY25 p.66", ""),
        ("Impact of derivatives hedging debt / fair value of hedges", {"30/06/2025": 211, "31/12/2025": 276, "30/06/2026": 244}, "H126 p.28; FY25 p.66", ""),
        ("Cash and cash equivalents", {"30/06/2025": -7330, "31/12/2025": -8021, "30/06/2026": -7288}, "H126 p.28; FY25 p.66", ""),
        ("Liquid assets and financing-related assets", {"30/06/2025": -2011, "31/12/2025": -1952, "30/06/2026": -2044}, "H126 p.28; FY25 p.66", ""),
        ("Remeasurement of financial liabilities (Suez PPA), excluded", {"30/06/2025": -187, "30/06/2026": -143}, "H126 p.28", "Not shown in the FY25 note: explains why 31/12/2025 is 19,821 there and 19,657 in the H1-26 bridge"),
        ("NET FINANCIAL DEBT", {"30/06/2025": 20764, "31/12/2025": 19821, "30/06/2026": 24548}, "H126 p.28; FY25 p.66", "31/12/2025 per FY25 note 8.3.2.1; the H1-26 bridge opens at 19,657"),
        ("Share of net debt at fixed rates after hedging (%)", {"30/06/2025": 90, "30/06/2026": 76}, "H126 p.28", ""),
        ("Average maturity of net debt (years)", {"30/06/2025": 6.6, "30/06/2026": 6.2}, "H126 p.28", ""),
    ]),
    ("Debt maturity at 31/12/2025 (undiscounted contractual flows, incl. interest)", ["2026", "2027", "2028", "2029", "2030", "Beyond 5 years", "Total"], [
        ("Bond issues", {"2026": 1924, "2027": 3043, "2028": 2447, "2029": 1781, "2030": 2006, "Beyond 5 years": 11065, "Total": 22265}, "FY25 p.66", ""),
        ("Other liabilities and bank overdrafts", {"2026": 7219, "2027": 465, "2028": 728, "2029": 302, "2030": 254, "Beyond 5 years": 1009, "Total": 9977}, "FY25 p.66", ""),
        ("IFRS 16 lease debt", {"2026": 499, "2027": 384, "2028": 317, "2029": 198, "2030": 147, "Beyond 5 years": 618, "Total": 2162}, "FY25 p.66", ""),
        ("Gross financial liabilities", {"2026": 9642, "2027": 3891, "2028": 3492, "2029": 2281, "2030": 2407, "Beyond 5 years": 12692, "Total": 34404}, "FY25 p.66", ""),
        ("Undrawn syndicated loan facility, by maturity", {"2030": 4500, "Total": 4500}, "FY25 p.67", "Extended to 2030 on 17/02/2025"),
        ("Undrawn credit lines, by maturity", {"2026": 227, "2027": 321, "2028": 590, "2029": 150, "2030": 150, "Total": 1438}, "FY25 p.67", ""),
    ]),
    ("Liquidity", ["31/12/2024", "31/12/2025", "30/06/2026"], [
        ("Undrawn syndicated loan facility (Veolia Environnement)", {"31/12/2024": 4500, "31/12/2025": 4500, "30/06/2026": 4500}, "FY25 p.66; H126 p.28", ""),
        ("Undrawn MT bilateral credit lines (Veolia Environnement)", {"31/12/2024": 724, "31/12/2025": 765, "30/06/2026": 765}, "FY25 p.66; H126 p.28", ""),
        ("Cash, liquid and financing assets (Veolia Environnement)", {"31/12/2024": 9349, "31/12/2025": 7551, "30/06/2026": 6767}, "FY25 p.66; H126 p.28", ""),
        ("Undrawn credit lines (subsidiaries)", {"31/12/2024": 949, "31/12/2025": 673, "30/06/2026": 608}, "FY25 p.66; H126 p.28", ""),
        ("Cash, liquid and financing assets (subsidiaries)", {"31/12/2024": 2270, "31/12/2025": 2422, "30/06/2026": 2565}, "FY25 p.66; H126 p.28", ""),
        ("TOTAL LIQUID ASSETS", {"31/12/2024": 17792, "31/12/2025": 15911, "30/06/2026": 15205}, "FY25 p.66; H126 p.28", ""),
        ("Current debt", {"31/12/2024": 9281, "31/12/2025": 8810, "30/06/2026": 11283}, "FY25 p.66; H126 p.28", ""),
        ("Bank overdrafts and other cash position items", {"31/12/2024": 197, "31/12/2025": 215, "30/06/2026": 90}, "FY25 p.66; H126 p.28", ""),
        ("LIQUIDITY NET OF CURRENT DEBT", {"31/12/2024": 8313, "31/12/2025": 6886, "30/06/2026": 3831}, "FY25 p.66; H126 p.28", ""),
    ]),
    ("Investments", ["H1 2024", "H1 2025", "H1 2026", "FY2023", "FY2024", "FY2025"], [
        ("Gross industrial investments", {"H1 2025": 1836, "H1 2026": 1746}, "H126 p.27", ""),
        ("  of which maintenance (incl. IFRS 16)", {"H1 2025": 842, "H1 2026": 757}, "H126 p.27", ""),
        ("  of which contractual growth", {"H1 2025": 737, "H1 2026": 704}, "H126 p.27", ""),
        ("  of which discretionary", {"H1 2025": 258, "H1 2026": 285}, "H126 p.27", ""),
        ("New operating financial assets", {"H1 2025": 106, "H1 2026": 112}, "H126 p.27", ""),
        ("Industrial divestitures", {"H1 2025": -89, "H1 2026": -116}, "H126 p.27", ""),
        ("Net industrial investments (incl. new operating financial assets)", {"H1 2024": 1722, "H1 2025": 1747, "H1 2026": 1630, "FY2023": 3730, "FY2024": 3836, "FY2025": 3855}, "H126 p.6, p.27; H124 p.6", ""),
        ("Financial acquisitions", {"H1 2024": 421, "H1 2025": 2168, "H1 2026": 2942}, "H124 p.29; H126 p.27", "Incl. acquisition costs and net debt of acquired entities"),
        ("Financial disposals", {"H1 2024": 253, "H1 2025": 18, "H1 2026": 73}, "H124 p.29; H126 p.27", ""),
        ("Net financial investments", {"H1 2024": 168, "H1 2025": 2150, "H1 2026": 2869}, "H124 p.29; H126 p.27", "Positive = net outflow"),
    ]),
    ("Acquisitions and disposals (EUR m)", ["Amount"], [
        ("Friedrich Hofmann (Germany, recycling) - 01/03/2024", {"Amount": 315}, "H124 p.29", ""),
        ("SADE (France) disposal - 29/02/2024", {"Amount": -175}, "H124 p.29", "Net proceeds"),
        ("Haikou (China) disposal - 26/06/2024", {"Amount": -79}, "H124 p.29", ""),
        ("Danubius (Hungary) - 06/01/2025", {"Amount": 271}, "FY25 p.17", "Net financial investment; securities price 366"),
        ("Water Technologies 30% minority from CDPQ - 30/06/2025", {"Amount": 1500}, "FY25 p.17", "USD 1.75bn; equity impact -1,380"),
        ("US hazardous waste tuck-ins + Chameleon - 2025", {"Amount": 247}, "FY25 p.17", "Net financial investment"),
        ("Zeeklite (Japan) - 30/05/2025", {"Amount": 85}, "FY25 p.17", ""),
        ("Enviropacific Services (Australia) - 31/03/2026", {"Amount": 137}, "H126 p.14, p.27", "AUD 228m"),
        ("Clean Earth (US) - price - 01/06/2026", {"Amount": 2542}, "H126 p.14", "USD 2,989m"),
        ("Clean Earth (US) - net financial investment incl. IFRS 16 debt and costs", {"Amount": 2778}, "H126 p.27", "Incl. 225 of IFRS 16 debt"),
        ("Belux stakes disposal - H1 2026", {"Amount": -41}, "H126 p.27", ""),
        ("France stakes disposal - H1 2026", {"Amount": -14}, "H126 p.27", ""),
    ]),
    ("Financing events, hybrids and rating", ["Amount"], [
        ("Bond issue 14/01/2026 (3 tranches, 2031/2034/2038, avg 3.583%)", {"Amount": 2500}, "H126 p.14", ""),
        ("Bond issue 10/04/2026 (2031 and 2036)", {"Amount": 1000}, "H126 p.14", ""),
        ("Non-dilutive convertible bond 29/06/2026 (0.75%, 2032)", {"Amount": 400}, "H126 p.14", ""),
        ("Bond repaid at maturity 09/06/2026", {"Amount": 750}, "H126 p.14", ""),
        ("Hybrid (deeply subordinated) redeemed 09/02/2026", {"Amount": 850}, "H126 p.14; FY25 p.74", "Reclassified to debt at 31/12/2025"),
        ("Hybrid debt outstanding at 31/12/2025 (before the redemption)", {"Amount": 4100}, "FY25 p.74", ""),
        ("Parent dividend paid 13/05/2026 (EUR 1.50 per share)", {"Amount": 1099}, "H126 p.14", ""),
        ("S&P rating: BBB / A-2, stable (confirmed 27/04/2026)", {}, "H126 p.14", "Text item"),
        ("Moody's rating: Baa1 / P-2, stable (confirmed 04/05/2026)", {}, "H126 p.14", "Text item"),
        ("Financial covenants in parent-company bank and bond documentation: none", {}, "FY25 p.67", "Text item"),
    ]),
    ("Guidance and targets", ["Value"], [
        ("2026: organic EBITDA growth, low end (%)", {"Value": 5}, "H126 p.29", ""),
        ("2026: organic EBITDA growth, high end (%)", {"Value": 6}, "H126 p.29", ""),
        ("2026: current net income growth, minimum (%)", {"Value": 8}, "H126 p.29", "At constant forex, incl. Clean Earth"),
        ("2026: net debt / EBITDA ('3x or slightly above')", {"Value": 3}, "H126 p.29", ""),
        ("Disposal programme by mid-2028 ('EUR 2bn+')", {"Value": 2000}, "H126 p.29", ""),
        ("H1 2026 FX effect on EBITDA (EUR m)", {"Value": -29}, "H126 p.22", "-0.9%"),
        ("GreenUp: EBITDA 2027 ('over EUR 8 billion')", {"Value": 8000}, "H124 p.12", ""),
        ("GreenUp: financial leverage ('less than or equal to 3x')", {"Value": 3}, "H124 p.12", ""),
        ("GreenUp: growth investments (EUR m), of which 2,000 for boosters", {"Value": 4000}, "H124 p.12", "Briefing p.20 says 'acquisitions over 2023-2027'"),
        ("GreenUp: annual savings (EUR m)", {"Value": 350}, "H124 p.12", ""),
        ("Clean Earth: EV / 2026e EBITDA post-synergies (x)", {"Value": 9.8}, "BRF p.21", ""),
        ("Clean Earth: run-rate synergies by year 4 (USD m)", {"Value": 120}, "BRF p.21", ""),
        ("Water Technologies 30%: EV / 2025 EBITDA post-synergies (x)", {"Value": 11}, "BRF p.21", ""),
        ("Water Technologies 30%: synergies by 2027 (EUR m)", {"Value": 90}, "BRF p.21; FY25 p.17", ""),
    ]),
]

# ------------------------------------------------------------------ 3. write the sheet
wb = Workbook()
rd = wb.active; rd.title = "Read me"
rd.column_dimensions["A"].width = 14; rd.column_dimensions["B"].width = 100
lines = [("Veolia - Step 1: input data", ""), ("", ""),
         ("What", "Every financial figure the capacity analysis may use, copied from the published documents. No calculation yet."),
         ("How", "Statements are read from the PDFs by extract_statements.py (no manual typing). Note figures are listed with their page."),
         ("Source codes", ""),
] + [(k, v) for k, v in DOCS.items()] + [("", ""),
         ("Signs", "As published. Statements: costs and outflows negative. Key figures: net debt and investments shown in brackets as in the report."),
         ("Colours", "Blue = number copied from a document. Black = formula (Checks block only)."),
         ("Page", "'p.' = PDF page number of the file (not the printed page number)."),
         ("Checks", "The Checks block at the bottom of the Inputs sheet re-adds each statement. Differences of 1-2 are rounding (the reports say so).")]
for i, (a, b) in enumerate(lines, 1):
    rd.cell(row=i, column=1, value=a).font = Font(name=F, bold=bool(a) and i > 1 or i == 1, size=14 if i == 1 else 10)
    rd.cell(row=i, column=2, value=b).font = Font(name=F, size=10)

ws = wb.create_sheet("Inputs")
ws.column_dimensions["A"].width = 62
for c in range(2, 10):
    ws.column_dimensions[L(c)].width = 12.5
ws.column_dimensions["J"].width = 24; ws.column_dimensions["K"].width = 60
ws["A1"] = "Veolia - input data (EUR m unless stated)"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "Blue = copied from the source document shown in column J. Nothing on this sheet is estimated."; ws["A2"].font = Font(name=F, italic=True, size=9, color=GREY)
ws.freeze_panes = "B4"
row = 4
CELL = {}   # (section, label, period) -> cell ref

def section_header(title, periods):
    global row
    ws.cell(row=row, column=1, value=title)
    for c in range(1, 12):
        ws.cell(row=row, column=c).fill = NAVY
        ws.cell(row=row, column=c).font = Font(name=F, bold=True, color="FFFFFF", size=10)
    for i, p in enumerate(periods):
        ws.cell(row=row, column=2 + i, value=p).alignment = Alignment(horizontal="right")
    ws.cell(row=row, column=10, value="Source"); ws.cell(row=row, column=11, value="Note")
    row += 1

def write_row(sec, label, periods, vals, src, note=""):
    global row
    big = label.isupper() or label.startswith(("Total", "TOTAL", "Net cash from", "NET"))
    c = ws.cell(row=row, column=1, value=label); c.font = Font(name=F, size=10, bold=big)
    for i, p in enumerate(periods):
        if p in vals:
            v = vals[p]
            cell = ws.cell(row=row, column=2 + i, value=v)
            cell.font = Font(name=F, size=10, color=BLUE, bold=big)
            cell.number_format = '0.00' if isinstance(v, float) and abs(v) < 20 and v != int(v) else NUM
            CELL[(sec, label, p)] = f"{L(2 + i)}{row}"
    ws.cell(row=row, column=10, value=src).font = Font(name=F, size=9, color=GREY)
    ws.cell(row=row, column=11, value=note).font = Font(name=F, size=9, color=GREY, italic=True)
    row += 1

for title, periods, parts in STATEMENTS:
    section_header(title, periods)
    for label, vals in merge(parts).items():
        srcs = sorted({s for _, s in vals.values()})
        write_row(title, label, periods, {p: v for p, (v, _) in vals.items()}, "; ".join(srcs))
    row += 1
for title, periods, items in NOTES:
    section_header(title, periods)
    for label, vals, src, note in items:
        write_row(title, label, periods, vals, src, note)
    row += 1

# ------------------------------------------------------------------ 4. checks (formulas)
section_header("Checks - should be 0 (or +/-2 rounding)", ["31/12/2024", "31/12/2025", "30/06/2026", "", "FY2024", "FY2025", "H1 2025", "H1 2026"])
A, EL, IS, CF, KF, ND = (STATEMENTS[1][0], STATEMENTS[2][0], STATEMENTS[0][0], STATEMENTS[3][0], STATEMENTS[4][0], NOTES[0][0])
def ref(sec, label, p):
    return CELL.get((sec, label, p))
def check(label, terms_by_col):
    """terms_by_col: {col_index: [(sign, sec, label, period)]}"""
    global row
    ws.cell(row=row, column=1, value=label).font = Font(name=F, size=10)
    for col, terms in terms_by_col.items():
        parts = []
        for sgn, sec, lab, p in terms:
            r = ref(sec, lab, p)
            if r is None:
                parts = None; break
            parts.append(f"{sgn}{r}")
        if parts:
            c = ws.cell(row=row, column=col, value="=" + "".join(parts)); c.font = Font(name=F, size=10); c.number_format = NUM
    row += 1

NCA = ["Goodwill", "Concession intangible assets", "Other intangible assets", "Property, plant and equipment", "Rights of use (net)", "Investments in joint ventures", "Investments in associates", "Non-consolidated investments", "Non-current operating financial assets", "Non-current derivative instruments - Assets", "Other non-current financial assets", "Deferred tax assets"]
CA = ["Inventories and work-in-progress", "Operating receivables", "Current operating financial assets", "Other current financial assets", "Current derivative instruments - Assets", "Cash and cash equivalents", "Assets classified as held for sale"]
NCL = ["Non-current provisions", "Non-current financial liabilities", "Non-current IFRS 16 lease debt", "Non-current derivative instruments - Liabilities", "Concession liabilities – non-current", "Deferred tax liabilities"]
CL = ["Operating payables", "Concession liabilities - current", "Current provisions", "Current financial liabilities", "Current IFRS 16 lease debt", "Current derivative instruments - Liabilities", "Bank overdrafts and other cash position items", "Liabilities directly associated with assets classified as held for sale"]
BS = {2: "31/12/2024", 3: "31/12/2025", 4: "30/06/2026"}
FL = {6: "FY2024", 7: "FY2025", 8: "H1 2025", 9: "H1 2026"}
def bs(total, items, sec):
    return {c: [("+", sec, total, p)] + [("-", sec, i, p) for i in items] for c, p in BS.items()}
check("Non-current assets = sum of lines", bs("Non-current assets", NCA, A))
check("Current assets = sum of lines", bs("Current assets", CA, A))
check("Total assets = non-current + current", bs("TOTAL ASSETS", ["Non-current assets", "Current assets"], A))
check("Non-current liabilities = sum of lines", bs("Non-current liabilities", NCL, EL))
check("Current liabilities = sum of lines", bs("Current liabilities", CL, EL))
check("Total equity and liabilities = equity + liabilities", bs("TOTAL EQUITY AND LIABILITIES", ["Equity", "Non-current liabilities", "Current liabilities"], EL))
check("Balance sheet balances (assets - equity and liabilities)", {c: [("+", A, "TOTAL ASSETS", p), ("-", EL, "TOTAL EQUITY AND LIABILITIES", p)] for c, p in BS.items()})
check("Income: pre-tax = operating income after equity-acc. + financial lines",
      {c: [("+", IS, "Pre-tax net income (loss)", p), ("-", IS, "Operating income after share of net income (loss) of equity-accounted entities", p), ("-", IS, "Cost of net financial debt", p), ("-", IS, "Other financial income and expenses", p)] for c, p in FL.items()})
check("Income: net income = continuing + discontinued",
      {c: [("+", IS, "Net income (loss) for the period", p), ("-", IS, "Net income (loss) from continuing operations", p), ("-", IS, "Net income (loss) from discontinued operations", p)] for c, p in FL.items()})
check("Cash flow: operating = before WCR + WCR + concessions WCR + taxes",
      {c: [("+", CF, "Net cash from operating activities of continuing operations", p), ("-", CF, "Operating cash flow before changes in working capital", p), ("-", CF, "Change in operating working capital requirements", p), ("-", CF, "Change in working capital requirements of concessions", p), ("-", CF, "Income taxes paid", p)] for c, p in FL.items()})
INV = ["Industrial investments, net of grants", "Proceeds on disposal of industrial assets", "Purchases of investments", "Proceeds on disposal of financial assets", "New operating financial assets", "Principal payments on operating financial assets", "Dividends received (including dividends received from joint ventures and associates)", "New non-current loans granted", "Principal payments on non-current loans", "Net decrease/increase in current loans"]
check("Cash flow: investing = sum of lines", {c: [("+", CF, "Net cash used in investing activities of continuing operations", p)] + [("-", CF, i, p) for i in INV] for c, p in FL.items()})
FIN = ["Net increase (decrease) in current financial liabilities", "Repayment of current IFRS 16 lease debt", "Other changes in non-current IFRS 16 lease debt", "New non-current borrowings and other debt", "Principal payments on non-current borrowings and other debt", "Change in liquid assets and financing financial assets", "Proceeds on issue of shares", "Share capital reduction", "Transactions with non-controlling interests: partial purchases", "Transactions with non-controlling interests: partial sales", "Issue / repayment of deeply subordinated securities", "Coupons on deeply subordinated securities", "Purchases of/proceeds from treasury shares", "Dividends paid", "Interest paid", "Interest on IFRIC 12 operating assets", "Interest on IFRS 16 lease debt"]
check("Cash flow: financing = sum of lines", {c: [("+", CF, "Net cash from (used in) financing activities of continuing operations", p)] + [("-", CF, i, p) for i in FIN] for c, p in FL.items()})
BEG = {"FY2024": "NET CASH AT THE BEGINNING OF THE YEAR", "FY2025": "NET CASH AT THE BEGINNING OF THE YEAR", "H1 2025": "NET CASH AT THE BEGINNING OF THE PERIOD", "H1 2026": "NET CASH AT THE BEGINNING OF THE PERIOD"}
END = {k: v.replace("BEGINNING", "END") for k, v in BEG.items()}
check("Cash flow: closing cash = opening + operating + investing + financing + FX",
      {c: [("+", CF, END[p], p), ("-", CF, BEG[p], p), ("-", CF, "Net cash from operating activities", p), ("-", CF, "Net cash used in investing activities", p), ("-", CF, "Net cash from (used in) financing activities", p), ("-", CF, "Effect of foreign exchange rate changes and other", p)] for c, p in FL.items()})
check("Net debt: composition 30/06/2026 adds up",
      {4: [("+", ND, "NET FINANCIAL DEBT", "30/06/2026"), ("-", ND, "Non-current financial liabilities", "30/06/2026"), ("-", ND, "Current financial liabilities", "30/06/2026"), ("-", ND, "Bank overdrafts and other cash position items", "30/06/2026"), ("-", ND, "Cash and cash equivalents", "30/06/2026"), ("-", ND, "Liquid assets and financing-related assets", "30/06/2026"), ("-", ND, "Impact of derivatives hedging debt / fair value of hedges", "30/06/2026"), ("-", ND, "Remeasurement of financial liabilities (Suez PPA), excluded", "30/06/2026")]})
check("Net debt: key figures 30/06/2026 vs composition (sign flipped)", {4: [("+", KF, "Net financial debt - Closing", "30/06/2026"), ("+", ND, "NET FINANCIAL DEBT", "30/06/2026")]})
check("Cash: balance sheet vs cash flow statement 30/06/2026", {4: [("+", A, "Cash and cash equivalents", "30/06/2026"), ("-", CF, "Cash and cash equivalents", "H1 2026")]})
for r_ in ws.iter_rows(min_row=1, max_row=row):
    for c in r_:
        if c.font and c.font.name != F:
            c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
ws.sheet_view.showGridLines = False
for wsx in wb.worksheets:
    wsx.page_setup.orientation = "landscape"; wsx.page_setup.paperSize = wsx.PAPERSIZE_A4
    wsx.page_setup.fitToWidth = 1; wsx.page_setup.fitToHeight = 0; wsx.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("saved", OUT, "rows", row, "cells", len(CELL))
